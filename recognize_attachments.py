#!/usr/bin/env python3
import os
import json
import re
import sqlite3
import zipfile
from pathlib import Path

# Optional imports
try:
    import fitz  # PyMuPDF
except Exception:
    fitz = None

try:
    import pdfplumber
except Exception:
    pdfplumber = None

try:
    import camelot
except Exception:
    camelot = None

try:
    import pytesseract
except Exception:
    pytesseract = None

try:
    from bs4 import BeautifulSoup
except Exception:
    BeautifulSoup = None

try:
    import lxml  # noqa: F401
except Exception:
    pass

try:
    import docx
except Exception:
    docx = None

try:
    import openpyxl
except Exception:
    openpyxl = None

DB_PATH = os.getenv('FOOD_DB_PATH', 'food.db')
ATTACHMENTS_ROOT = os.getenv('LOCAL_ATTACHMENTS_ROOT', 'attachments')
PARSED_ROOT = os.getenv('PARSED_ROOT', 'parsed')


def init_db():
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute('''
        CREATE TABLE IF NOT EXISTS parsed_documents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            reestr_number TEXT,
            file_name TEXT,
            file_path TEXT,
            parsed_path TEXT,
            file_type TEXT,
            status TEXT,
            created_at TEXT
        )
    ''')
    cur.execute('''
        CREATE TABLE IF NOT EXISTS contract_types (
            reestr_number TEXT PRIMARY KEY,
            contract_type TEXT,
            source TEXT,
            confidence TEXT,
            created_at TEXT
        )
    ''')
    conn.commit()
    conn.close()


def save_parsed_meta(meta):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO parsed_documents (reestr_number, file_name, file_path, parsed_path, file_type, status, created_at)
        VALUES (?, ?, ?, ?, ?, ?, datetime('now'))
    ''', (
        meta.get('reestr_number'),
        meta.get('file_name'),
        meta.get('file_path'),
        meta.get('parsed_path'),
        meta.get('file_type'),
        meta.get('status')
    ))
    conn.commit()
    conn.close()


def write_json(path, data):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def extract_pdf_text(path):
    text = ""
    if pdfplumber:
        with pdfplumber.open(path) as pdf:
            for page in pdf.pages:
                text += (page.extract_text() or "") + "\n"
    elif fitz:
        doc = fitz.open(path)
        for page in doc:
            text += page.get_text() + "\n"
    return text.strip()


def extract_pdf_tables(path):
    if not camelot:
        return []
    try:
        tables = camelot.read_pdf(path, pages='all')
        return [t.df.values.tolist() for t in tables]
    except Exception:
        return []


def extract_docx_text(path):
    if not docx:
        return ""
    doc = docx.Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def extract_xlsx(path):
    if not openpyxl:
        return {}
    wb = openpyxl.load_workbook(path, data_only=True)
    data = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else v for v in row])
        data[ws.title] = rows
    return data


def extract_html_text(path):
    raw = Path(path).read_text(encoding='utf-8', errors='ignore')
    if not BeautifulSoup:
        return raw
    soup = BeautifulSoup(raw, 'html.parser')
    return soup.get_text("\n").strip()


def extract_xml_text(path):
    raw = Path(path).read_text(encoding='utf-8', errors='ignore')
    # Strip tags crudely for now
    text = re.sub(r"<[^>]+>", " ", raw)
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def classify_from_web_text(text):
    if not text:
        return None, "web", "low"
    lower = text.lower()
    keywords = [
        "питание", "завтрак", "обед", "ужин", "дет", "школ",
        "столов", "пищ", "рацион", "продукт"
    ]
    if any(k in lower for k in keywords):
        return "питание", "web", "high"
    return None, "web", "low"


def extract_object_block(html_text):
    idx = html_text.lower().find("объекты закупки")
    if idx == -1:
        return ""
    return html_text[idx:idx + 2000]


def save_contract_type(reestr_number, contract_type, source, confidence):
    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    cur.execute('''
        INSERT INTO contract_types (reestr_number, contract_type, source, confidence, created_at)
        VALUES (?, ?, ?, ?, datetime('now'))
        ON CONFLICT(reestr_number) DO UPDATE SET
            contract_type=excluded.contract_type,
            source=excluded.source,
            confidence=excluded.confidence,
            created_at=excluded.created_at
    ''', (reestr_number, contract_type, source, confidence))
    conn.commit()
    conn.close()


def handle_file(reestr_number, file_path):
    ext = file_path.suffix.lower()
    parsed = {
        'file': file_path.name,
        'type': ext,
        'text': '',
        'tables': []
    }

    status = 'parsed'
    try:
        if ext in ['.pdf']:
            parsed['text'] = extract_pdf_text(str(file_path))
            parsed['tables'] = extract_pdf_tables(str(file_path))
        elif ext in ['.doc', '.docx']:
            parsed['text'] = extract_docx_text(str(file_path))
        elif ext in ['.xls', '.xlsx']:
            parsed['tables'] = extract_xlsx(str(file_path))
        elif ext in ['.html', '.htm']:
            parsed['text'] = extract_html_text(str(file_path))
        elif ext in ['.xml']:
            parsed['text'] = extract_xml_text(str(file_path))
        else:
            status = 'skipped'
    except Exception as e:
        status = f'error: {e}'

    parsed_dir = Path(PARSED_ROOT) / reestr_number
    parsed_path = parsed_dir / f"{file_path.name}.json"
    write_json(parsed_path, parsed)

    save_parsed_meta({
        'reestr_number': reestr_number,
        'file_name': file_path.name,
        'file_path': str(file_path),
        'parsed_path': str(parsed_path),
        'file_type': ext,
        'status': status
    })


def scan_contract(reestr_number):
    root = Path(ATTACHMENTS_ROOT) / reestr_number
    if not root.exists():
        print(f"No attachments for {reestr_number}")
        return

    # Unzip any zip files
    for zip_path in root.glob('*.zip'):
        try:
            with zipfile.ZipFile(zip_path, 'r') as z:
                z.extractall(root / zip_path.stem)
        except Exception:
            pass

    for file_path in root.rglob('*'):
        if file_path.is_file():
            handle_file(reestr_number, file_path)

    # Classify contract type from web page (document-info.html)
    web_html = root / 'document-info.html'
    if web_html.exists():
        html_text = web_html.read_text(encoding='utf-8', errors='ignore')
        block = extract_object_block(html_text)
        ctype, source, confidence = classify_from_web_text(block)
        if not ctype:
            try:
                from llm_gateway import classify_contract_type
                ctype = classify_contract_type(block or html_text)
                source = "llm"
                confidence = "medium"
            except Exception:
                ctype = "прочее"
                source = "fallback"
                confidence = "low"
        save_contract_type(reestr_number, ctype, source, confidence)


def main():
    import sys
    if len(sys.argv) < 2:
        print("Usage: recognize_attachments.py <reestr_number>")
        return

    reestr_number = sys.argv[1]
    init_db()
    scan_contract(reestr_number)
    print("Done")


if __name__ == '__main__':
    main()
